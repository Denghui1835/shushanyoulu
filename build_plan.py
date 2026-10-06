# -*- coding: utf-8 -*-
"""生成《书山有路——课程的 GitHub 与听读一体的开源自学习平台规划书》，照公文格式。
规格沿用 D:\\SRT项目\\_reformat_gongwen.py（用户 2026-10-05 定死的那套）。
内容源：同目录 PRD.md v0.2，加上 2026-10-06 定的「免费开源 + 算力自带 + 班级」三条口径。
改内容只动下面的 C 列表，别动渲染引擎。"""
import os, io, sys, re, zipfile
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')

OUT_DIR = r'D:\AI coding\书山有路'
DST = os.path.join(OUT_DIR, '书山有路_产品规划书_v1.docx')

FS, HT, KT, HZ = '仿宋', '黑体', '楷体', '华文中宋'
SZ = 16          # 三号
LS = 28          # 行距固定值 28 磅
TBLZ = 14        # 表格四号
TBL_LS = 22
WEST = 'Times New Roman'
TEXTW = Cm(15.6)

# ============================================================
# 内容定义
#   ('title'|'subtitle'|'info'|'h0'|'h1'|'h2'|'p'|'li'|'caption'|'table'|'code'|'blank'|'pagebreak', payload)
# ============================================================
C = []
A = C.append

# ---------- 封面 ----------
A(('title', '书山有路'))
A(('subtitle', '——课程的 GitHub 与听读一体的自学习平台规划书'))
A(('info', '文档版本：v1.0'))
A(('info', '编制日期：2026年10月6日'))
A(('info', '文档性质：产品规划书（招募共建者）'))
A(('blank', ''))

# ---------- 一 ----------
A(('h0', '一、这个产品是什么'))

A(('h1', '（一）一句话：课程的 GitHub'))
A(('p', '书山 Hub 是课程的 GitHub。人人可以用它做一门自己喜欢的课，上传托管；愿意开源就开源，不愿意就私有；想找人一起改，就拉协作者，把一门课改到臻于至善。'))
A(('p', '这里的“课”不是一个抽象说法，而是一个能装下完整内容的结构：一本教材或一份讲义，拆成章；每一章下面挂着原文、知识点、题目、讲稿、音频。它既能读，也能听，也能刷题。'))

A(('h1', '（二）另一句话：听读一体，听带读'))
A(('p', '走路、通勤、睡前都能学。**听是主线，但听的时候能跟着原文走**——不至于听个热闹。'))
A(('p', '这一条是整个产品存在的理由。只听不跟原文，等于没学；只读不借助耳朵，又太费劲。我们要的就是那个交集。'))

A(('h1', '（三）为什么是这个交集'))
A(('table', [
    ['学习方式', '爽（不痛苦）', '真正受益'],
    ['抖音式短视频', '是', '否：刷完就忘'],
    ['课本', '否：读起来痛苦', '是'],
    ['书山有路', '是：听着不痛苦', '是：跟着原文，学得进去'],
]))
A(('p', '我们不做好坏判断，只承认一个事实：**能被坚持的学习方式才有用**。短视频的问题不是短，是注意力离开了要学的东西；课本的问题不是深，是读起来太费劲。'))

# ---------- 二 ----------
A(('h0', '二、三条产品线（并列，一条不删）'))

A(('h1', '（一）备考线：已经建成的那条'))
A(('p', '计算机二级 Python 与 C 语言两门预置课：知识树、题库、刷题、全真模拟、闪卡、关键词挖空背诵。这一条已经建成并在用，**一行不删**。'))
A(('p', '它是整个产品的入口。人是因为要考试才第一次打开它，然后才可能留下来听一门课。'))
A(('h1', '（二）听读线：新的主线'))
A(('p', '原文逐句朗读、逐句高亮、点句跳转；后续加入单人口语讲书、双主播对谈。这是本次建设的重心。'))

A(('h1', '（三）书山 Hub：内容与协作层'))
A(('p', '广场、发布、收藏、一键学习、Fork、建议这些骨架已经存在；本轮补上可见性、许可、协作者与版本快照。'))

A(('p', '三条线的关系：Hub 里发布的仓库，其章节自动成为别人“听读”页上可听的专辑；平台预置的精品课，就是官方账号名下公开的仓库，音频全局只合成一次，N 个人共享。'))

# ---------- 三 ----------
A(('h0', '三、听读的三种形态：用户自选'))

A(('table', [
    ['形态', '音频内容', '与原文对齐', '需要模型', '定位'],
    ['read', '原文逐句朗读，一字不改', '逐句高亮，来自分句本身', '否', '最忠实，先落地'],
    ['narrate', '讲书人口语化改写（段级）', '模型插入锚点，失败时模糊回退', '是', '主线默认'],
    ['dialogue', '双主播对谈（复用现有播客）', '段级加锚点', '少量（一次调用）', '节目化副产品'],
]))

A(('p', '这里有一个要说清楚的取舍：双主播对谈是“另一个人替你消化”——听着爽，但**注意力离开了原文**，与“能跟着原文”天然冲突。所以单人讲书才最容易做到听是主线、读跟着走。'))
A(('p', '落地的顺序是：read 先落地，narrate 做成主线默认，dialogue 作为节目线保留（现有播客能力不浪费）。'))

# ---------- 四 ----------
A(('h0', '四、班级与「一起学」'))

A(('p', '一个人学，最容易发生的事就是今天算了。'))

A(('h1', '（一）已经有了的：一起学'))
A(('p', '学习行为（打卡、完成任务、学完一讲、模拟考）会自动生成一条动态，别人可以点赞鼓励，每周有一个学习榜。这部分已经建成了。'))

A(('h1', '（二）班级：把大广场收成一个小圈子'))
A(('p', '现在的动态流是全站所有人的。**班级是在它上面加一层“只看我们班”**。一个班有班长和成员，有公告、有共同在学的课、有班内的榜。'))
A(('p', '一个班可以挂多门课，一门课也可以被多个班用。Hub 是公开的课表，班级是把人聚到课上的圈子。'))

A(('h1', '（三）和已有工具的区别'))
A(('p', '学习通是**管理工具**：签到、点名、发作业，动力来自考核。我们做的是**一起学**：动力来自结伴——同一个班的人看得见彼此今天的进度，谁掉队了看得见，可以被拉一把。这一条不靠强制，靠的是熟人之间那点不好意思。'))
A(('p', '班级排在路线图的 P3.5。**先要有东西给人学，人才聚得起来**。'))

# ---------- 五 ----------
A(('h0', '五、免费与开源：这件事怎么活下去'))

A(('h1', '（一）我们的承诺'))
A(('p', '**所有课程内容免费听、免费学；代码开源。**不设会员，不设付费墙。任何人把自己的课放上来，别人就能直接听，不用先付钱。'))

A(('h1', '（二）算力自带：我们只提供框架'))
A(('p', '这是本方案最关键的一条约定。**我们只开发框架，不承担模型算力。**具体是：'))
A(('li', '用户可以配置自己的模型接口，配好之后所有智能功能走他自己的额度，不消耗平台的；'))
A(('li', '也可以完全本地运行，数据留在自己电脑上——本项目本来就是本地优先的单库结构，天生适合各自装一份；'))
A(('li', '平台只负责把“怎么用”做顺，不负责替谁付钱。'))
A(('p', '这条约定同时解决两个问题。一是**可持续**：没有无限补贴的窟窿，项目不会因为有人听得多就垮掉。二是**可信**：数据在你自己机器上，谁也不担心自己的学习记录被拿去做什么。'))

A(('h1', '（三）预置精品课：官方兜底'))
A(('p', '冷启动需要内容，所以官方出一门预置精品课。音频按内容寻址缓存，**全局只合成一次、N 个人听同一份**，成本不随人数线性上涨。官方只出这一份，剩下的靠作者。'))

A(('h1', '（四）我们真正出的是什么'))
A(('p', '**我们不出算力，我们出的是那套让别人愿意把课放上来、愿意一起改、改了别人还能直接听的框架。**这才是这个项目要做的东西。'))

# ---------- 六 ----------
A(('h0', '六、一起干：缺哪些人，能认领哪些活'))

A(('p', '这份规划书是写给志同道合的同学看的。下面这些活现在缺人，可以认领：'))
A(('table', [
    ['可以认领的活', '要做什么', '需要什么'],
    ['写课', '把你真正懂的那门课，整理成一门能听、能读、能刷的课', '懂那门课就行'],
    ['前端', '听读页、班级页、Hub 的交互与视觉', 'React 基础'],
    ['后端', '听读管线、协作与版本、权限', 'Python 与 FastAPI 基础'],
    ['移动端', 'Expo 客户端，让通勤路上真的能用上', 'React Native 基础'],
    ['班长', '在自己的班里带一群人学下去', '愿意管人'],
    ['内容与版权', '上传内容的边界、预置课的授权来源', '细心'],
]))
A(('p', '我负责整体的产品方向，以及把框架搭起来。**你不需要是全栈，认领一块就够。**'))

# ---------- 七 ----------
A(('h0', '七、技术怎么落地'))

A(('h1', '（一）整体架构'))
A(('code', '前端    React + Vite + Ant Design (:5173)   ｜   Expo / RN 移动端'))
A(('code', '                              ↓  HTTP / SSE'))
A(('code', '服务    FastAPI + SQLAlchemy (async)   (:8000)'))
A(('code', '                              ↓'))
A(('code', '数据    SQLite（本地优先，WAL）'))
A(('code', '                              ↓'))
A(('code', '模型    LLM：用户自带接口（DeepSeek / Claude / OpenAI / Ollama）'))
A(('code', '        TTS：edge-tts（默认，免费免 Key）/ 豆包 / Azure'))

A(('h1', '（二）听读管线'))
A(('code', 'Document → 取阅读单元 → 拼整章文本'))
A(('code', '  ├ read    ：分句 → 分组（每段不超过 180 字）→ 逐段合成 → 实测时长 → 逐句时间轴'))
A(('code', '  ├ narrate ：按块切分（约 3500 字）→ 每块一次模型调用（口语讲稿 + 锚点）→ 拼装'))
A(('code', '  └ dialogue：生成播客文稿 → 解析轮次 → 合并 → 合成'))

A(('h1', '（三）三级复用：预置精品课的命门'))
A(('li', '同一仓库同一讲重新生成时命中文稿指纹，直接复用文稿；'))
A(('li', '**全局音频去重**：音频按内容寻址存放，预置课只生成一次、N 人共享；'))
A(('li', '跨仓库相同源文本命中文稿指纹，连模型也只跑一次。'))
A(('p', '这三条是“预置精品课只花一次钱”的技术保证，也是免费模式能站住的前提。'))

A(('h1', '（四）其他关键设计'))
A(('li', '轻量检索：用中文分词做词法检索，替代重型向量库，接口独立，随时可换；'))
A(('li', '复用既有的模型适配层：缓存、限流、重试、多供应商统一处理；'))
A(('li', '本地优先：单个 SQLite 库，启动即用，不需要部署数据库。'))

# ---------- 八 ----------
A(('h0', '八、关键决策记录'))

A(('table', [
    ['序号', '决策', '结论', '理由'],
    ['1', '核心体验', '听读为主线', '通勤、走路、睡前都能学'],
    ['2', '收听形态', '听读一体：听带读', '只听会走神，不听原文等于没学'],
    ['3', '现有备考模块', '并列保留，一行不删', '已建成且有人用，拆掉是纯损失'],
    ['4', '产品性质', '公开产品，免费使用', '想让大家都用得上'],
    ['5', '成本模式', '内容免费、算力自带、代码开源', '不做无限补贴，才活得下去'],
    ['6', '内容来源', '平台预置精品课，加用户导入', '冷启动要内容，长期靠作者'],
    ['7', '命名', '应用叫书山有路，社区叫书山 Hub', '保留既有品牌，新概念单独命名'],
    ['8', '听读形态', '三种都留，用户自选', '不同场景要不同的东西'],
    ['9', '时间轴方案', '逐段实测优先，不引新依赖、不用 SSML', '与供应商无关，段间不累积漂移'],
    ['10', 'Hub 协作', '邀请编辑先做，变更集后做，真 git 不做', '先让“邀请同学一起改”跑起来'],
    ['11', '班级', '排在 P3.5', '先有内容，人才聚得起来'],
]))

# ---------- 九 ----------
A(('h0', '九、路线图'))

A(('table', [
    ['阶段', '内容', '状态'],
    ['P0 底线修复', '登录注册界面、按用户过滤、越权防护', '已完成'],
    ['P1 纵向切片', '原文朗读、逐句时间轴、内容寻址缓存、听读页逐句高亮', '进行中'],
    ['P2 分讲化', '讲与音频资产入库、异步生成与进度推送、连播、断点续听', '待做'],
    ['P3 讲书人模式', '口语化改写、源文跟随高亮', '待做'],
    ['P3.5 班级', '建班、成员、班内公告与共同课程、班内榜', '待做'],
    ['P4 导航重构', '四组菜单：听读、备考、书山 Hub、我的', '待做'],
    ['P5 Hub 基础', '可见性、许可、协作者、发布入口、版本快照', '待做'],
    ['P6 Hub 协作流', '带变更集的建议、合并、协作者管理', '待做'],
    ['P7 多端打磨', '移动端后台播放、连播、进度同步', '待做'],
]))

# ---------- 十 ----------
A(('h0', '十、风险与未决'))

A(('table', [
    ['事项', '说明'],
    ['首发切入口', '先服务身边能直接邀请到的人，还是直接打全网，尚未最终确定。Hub 要内容供给，而作者只可能来自你能直接邀请到的人'],
    ['内容版权', '用户上传他人教材的边界，以及预置精品课的授权来源，都需要明确'],
    ['配音成本与限流', '免费线路有限流，整章连播需要异步化才撑得住'],
    ['协作可行性', '“变更集加快照”能否覆盖真实协作场景，需要在 P6 用两个真人协作验证'],
    ['移动端后台播放', '通勤场景的命门，但排在 P7，偏晚'],
]))

# ---------- 十一 ----------
A(('h0', '十一、怎么跑起来'))

A(('code', '# 前置：Python 3.10 以上、Node 18 以上'))
A(('code', '# 复制 backend/.env.example 为 .env，填入你自己的模型接口'))
A(('code', 'start.bat'))
A(('p', '后端默认监听 8000，前端默认 5173，启动即用。'))

# ============================================================
# 渲染
# ============================================================
doc = Document()

# --- 页面设置 ---
for s in doc.sections:
    s.top_margin, s.bottom_margin = Cm(3.7), Cm(3.5)
    s.left_margin, s.right_margin = Cm(2.8), Cm(2.6)

# --- Normal 基线 ---
n = doc.styles['Normal']
n.font.size = Pt(SZ)
rpr = n.element.get_or_add_rPr()
rf = rpr.get_or_add_rFonts()
rf.set(qn('w:eastAsia'), FS); rf.set(qn('w:ascii'), WEST); rf.set(qn('w:hAnsi'), WEST)

def set_run(run, ea, size=SZ, bold=False):
    run.font.size = Pt(size)
    run.bold = bold
    rpr = run._element.get_or_add_rPr()
    rf = rpr.get_or_add_rFonts()
    rf.set(qn('w:eastAsia'), ea)
    rf.set(qn('w:ascii'), WEST); rf.set(qn('w:hAnsi'), WEST)

def set_ind(p, first_chars=None, left=None, hanging=None):
    pPr = p._p.get_or_add_pPr()
    ind = pPr.find(qn('w:ind'))
    if ind is None:
        ind = OxmlElement('w:ind'); pPr.append(ind)
    for k in ('w:firstLine', 'w:firstLineChars', 'w:left', 'w:leftChars',
              'w:hanging', 'w:hangingChars'):
        if ind.get(qn(k)) is not None:
            del ind.attrib[qn(k)]
    if first_chars is not None:
        ind.set(qn('w:firstLineChars'), str(first_chars))
        ind.set(qn('w:firstLine'), str(int(first_chars * 3.2)))
    if left is not None:
        ind.set(qn('w:left'), str(left)); ind.set(qn('w:leftChars'), '0')
    if hanging is not None:
        ind.set(qn('w:hanging'), str(hanging)); ind.set(qn('w:hangingChars'), '0')

def set_line(p, pt=LS, before=0, after=0):
    pf = p.paragraph_format
    pf.line_spacing = Pt(pt)
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)

def set_outline(p, lvl):
    pPr = p._p.get_or_add_pPr()
    old = pPr.find(qn('w:outlineLvl'))
    if old is not None:
        pPr.remove(old)
    e = OxmlElement('w:outlineLvl')
    e.set(qn('w:val'), str(lvl))
    rpr = pPr.find(qn('w:rPr'))
    if rpr is not None:
        rpr.addprevious(e)
    else:
        pPr.append(e)

def set_page_break_before(p):
    pPr = p._p.get_or_add_pPr()
    old = pPr.find(qn('w:pageBreakBefore'))
    if old is not None:
        pPr.remove(old)
    e = OxmlElement('w:pageBreakBefore')
    st = pPr.find(qn('w:pStyle'))
    if st is not None:
        st.addnext(e)
    else:
        pPr.insert(0, e)

BOLD_RE = re.compile(r'\*\*(.+?)\*\*')
def emit_runs(p, text, ea, size=SZ, bold_all=False):
    """支持 **加粗** 内联标记。"""
    pos = 0
    for m in BOLD_RE.finditer(text):
        if m.start() > pos:
            set_run(p.add_run(text[pos:m.start()]), ea, size=size, bold=bold_all)
        set_run(p.add_run(m.group(1)), ea, size=size, bold=True)
        pos = m.end()
    if pos < len(text):
        set_run(p.add_run(text[pos:]), ea, size=size, bold=bold_all)
    if not text:
        set_run(p.add_run(''), ea, size=size, bold=bold_all)

# --- 列宽：按内容分配（短列窄、长列宽，绝不均分）---
def col_widths(rows):
    ncol = len(rows[0])
    weights = []
    for ci in range(ncol):
        mx = 0
        for r in rows:
            if ci >= len(r):
                continue
            cell = r[ci]
            w = 0.0
            for ch in cell:
                w += 1.0 if ord(ch) > 0x2E80 else 0.55
            mx = max(mx, w)
        weights.append(max(mx, 2.0))
    total = sum(weights)
    MIN = 1.6   # cm
    widths = [TEXTW * (w / total) for w in weights]
    # 抬升过窄列，等比压缩其余
    for _ in range(3):
        deficit = sum(max(0.0, MIN - w) for w in widths)
        if deficit < 0.01:
            break
        widths = [w if w >= MIN else MIN for w in widths]
        over = sum(w for w in widths if w > MIN)
        target = float(TEXTW) - MIN * sum(1 for w in widths if w <= MIN)
        if over > 0:
            widths = [w if w <= MIN else w * (target / over) for w in widths]
    s = sum(widths)
    return [int(float(TEXTW) * (w / s)) for w in widths]

def add_table(rows):
    t = doc.add_table(rows=len(rows), cols=len(rows[0]))
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    ws = col_widths(rows)
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            cell = t.cell(ri, ci)
            cell.width = ws[ci]
            cp = cell.paragraphs[0]
            for r in list(cp.runs):
                r._element.getparent().remove(r._element)
            set_run(cp.add_run(val), HT if ri == 0 else FS,
                    size=TBLZ, bold=(ri == 0))
            set_ind(cp, first_chars=0)
            set_line(cp, TBL_LS)
    return t

def add_toc_field(p):
    def mkfld(kind, dirty=False):
        r = p.add_run()
        fc = OxmlElement('w:fldChar')
        fc.set(qn('w:fldCharType'), kind)
        if dirty:
            fc.set(qn('w:dirty'), 'true')
        r._r.append(fc)
    mkfld('begin', dirty=True)
    r = p.add_run()
    it = OxmlElement('w:instrText')
    it.set(qn('xml:space'), 'preserve')
    it.text = ' TOC \\o "1-2" \\h \\z \\u '
    r._r.append(it)
    mkfld('separate')
    p.add_run('（在 WPS 中按 Ctrl+A 后按 F9 更新域，生成目录）')
    mkfld('end')
    for r in p.runs:
        set_run(r, FS)

# --- 主渲染 ---
li_no = 0
toc_anchor = None
title_done = False

for kind, payload in C:
    if kind == 'title':
        p = doc.add_paragraph(); emit_runs(p, payload, HZ, size=22)
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_ind(p, first_chars=0); set_line(p, 34, after=6)
        title_done = True
        continue
    if kind == 'subtitle':
        p = doc.add_paragraph(); emit_runs(p, payload, KT)
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_ind(p, first_chars=0); set_line(p, LS, after=18)
        continue
    if kind == 'info':
        p = doc.add_paragraph(); emit_runs(p, payload, FS)
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_ind(p, first_chars=0); set_line(p, LS)
        continue
    if kind == 'blank':
        p = doc.add_paragraph(); set_ind(p, first_chars=0); set_line(p, LS)
        toc_anchor = p
        continue
    if kind == 'pagebreak':
        p = doc.add_paragraph(); set_ind(p, first_chars=0); set_line(p, LS)
        set_page_break_before(p)
        continue
    if kind == 'h0':
        li_no = 0
        p = doc.add_paragraph(); emit_runs(p, payload, HT)
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_ind(p, first_chars=0); set_line(p, LS, before=14, after=8)
        set_outline(p, 0)
        continue
    if kind == 'h1':
        li_no = 0
        p = doc.add_paragraph(); emit_runs(p, payload, HT)
        set_ind(p, first_chars=200); set_line(p, LS, before=8, after=4)
        set_outline(p, 1)
        continue
    if kind == 'h2':
        li_no = 0
        p = doc.add_paragraph(); emit_runs(p, payload, KT)
        set_ind(p, first_chars=200); set_line(p, LS, before=6, after=2)
        set_outline(p, 2)
        continue
    if kind == 'p':
        p = doc.add_paragraph(); emit_runs(p, payload, FS)
        set_ind(p, first_chars=200); set_line(p, LS)
        continue
    if kind == 'li':
        li_no += 1
        p = doc.add_paragraph()
        emit_runs(p, '%d. %s' % (li_no, payload), FS)
        set_ind(p, first_chars=0, left=640, hanging=640)
        set_line(p, LS)
        continue
    if kind == 'code':
        p = doc.add_paragraph(); emit_runs(p, payload, FS, size=12)
        set_ind(p, first_chars=0, left=420); set_line(p, 16)
        continue
    if kind == 'caption':
        li_no = 0
        p = doc.add_paragraph(); emit_runs(p, payload, HT, size=TBLZ)
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_ind(p, first_chars=0); set_line(p, TBL_LS, before=8, after=3)
        continue
    if kind == 'table':
        add_table(payload)
        p = doc.add_paragraph(); set_ind(p, first_chars=0); set_line(p, 10)
        continue

# --- 页脚页码「— 1 —」---
def page_footer(section):
    section.footer.is_linked_to_previous = False
    p = section.footer.paragraphs[0]
    for r in list(p.runs):
        r._element.getparent().remove(r._element)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    def mk(txt=None, field=False):
        r = p.add_run(txt or '')
        set_run(r, '宋体', size=14)
        if field:
            a = OxmlElement('w:fldChar'); a.set(qn('w:fldCharType'), 'begin')
            b = OxmlElement('w:instrText'); b.set(qn('xml:space'), 'preserve'); b.text = ' PAGE '
            c = OxmlElement('w:fldChar'); c.set(qn('w:fldCharType'), 'end')
            r._r.append(a); r._r.append(b); r._r.append(c)
        return r
    mk('— '); mk(field=True); mk(' —')
    set_line(p, 14)

for s in doc.sections:
    page_footer(s)

# --- 目录（放在正文前）---
if toc_anchor is not None:
    p_title = doc.add_paragraph()
    p_title.add_run('目　录')
    p_title.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_ind(p_title, first_chars=0); set_line(p_title, 34, after=12)
    for r in p_title.runs:
        set_run(r, HZ, size=22)
    set_page_break_before(p_title)

    p_toc = doc.add_paragraph()
    add_toc_field(p_toc)
    set_ind(p_toc, first_chars=0); set_line(p_toc, LS)

    anchor_el = toc_anchor._p
    for el in (p_toc._p, p_title._p):
        el.getparent().remove(el)
        anchor_el.addnext(el)

    # 正文首页（第一个 h0）另起一页
    for p in doc.paragraphs:
        if p.text.strip().startswith('一、这个产品'):
            set_page_break_before(p)
            break

# --- 收尾：抹掉样式名，统一字体 ---
for p in doc.paragraphs:
    if p.style.name != 'Normal':
        p.style = doc.styles['Normal']
for p in doc.paragraphs:
    for r in p.runs:
        rpr = r._element.get_or_add_rPr()
        rf = rpr.find(qn('w:rFonts'))
        if rf is None:
            rf = OxmlElement('w:rFonts'); rpr.insert(0, rf)
        if rf.get(qn('w:eastAsia')) == '微软雅黑':
            rf.set(qn('w:eastAsia'), FS)

os.makedirs(OUT_DIR, exist_ok=True)
doc.save(DST)

# --- 清掉模板带的孤立字体引用（WPS「缺失字体」告警的两个来源）---
def patch_zip(path):
    tmp = path + '.tmp'
    zin = zipfile.ZipFile(path, 'r')
    zout = zipfile.ZipFile(tmp, 'w', zipfile.ZIP_DEFLATED)
    for it in zin.infolist():
        data = zin.read(it.filename)
        name = it.filename
        if name in ('word/styles.xml', 'word/stylesWithEffects.xml'):
            s = data.decode('utf-8')
            for f in ('Courier',):
                s = s.replace('w:ascii="%s"' % f, 'w:ascii="Times New Roman"')
                s = s.replace('w:hAnsi="%s"' % f, 'w:hAnsi="Times New Roman"')
                s = s.replace('w:eastAsia="%s"' % f, 'w:eastAsia="Times New Roman"')
                s = s.replace('w:cs="%s"' % f, 'w:cs="Times New Roman"')
            s = s.replace('"微软雅黑"', '"仿宋"')
            data = s.encode('utf-8')
        elif name == 'word/theme/theme1.xml':
            s = data.decode('utf-8')
            # 脚本字体表里 26 个本机没有的字体（泰文/天城文等）→ 整表清空
            s = re.sub(r'<a:font script="[^"]*" typeface="[^"]*"\s*/>', '', s)
            data = s.encode('utf-8')
        zout.writestr(it, data)
    zin.close(); zout.close()
    os.replace(tmp, path)

patch_zip(DST)

print('OK ->', DST)
print('段落', len(doc.paragraphs), '表格', len(doc.tables))
